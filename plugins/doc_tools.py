"""Document handling tools — PDF, DOCX, CSV with secure guards."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from core.tools import Tool

_MAX_FILE_BYTES = 10 * 1024 * 1024
_ALLOWED_EXTS = {".pdf", ".docx", ".csv", ".txt", ".md", ".json"}

def _guard_path(path: str) -> tuple[Path | None, str | None]:
    if ".." in Path(path).parts:
        return None, "Error: path may not contain '..'"
    p = Path(path).expanduser().resolve()
    if not p.exists():
        return None, f"Error: file not found: {p}"
    if p.suffix.lower() not in _ALLOWED_EXTS:
        return None, f"Error: unsupported type {p.suffix} (allowed: {', '.join(sorted(_ALLOWED_EXTS))})"
    if p.stat().st_size > _MAX_FILE_BYTES:
        return None, f"Error: file too large ({p.stat().st_size} bytes, max {_MAX_FILE_BYTES})"
    return p, None

def pdf_read(path: str, max_pages: int = 20) -> str:
    p, err = _guard_path(path)
    if err:
        return err
    if p.suffix.lower() != ".pdf":
        return "Error: not a PDF file."
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(p))
        max_pages = max(1, min(int(max_pages), 50))
        texts = []
        for i in range(min(len(doc), max_pages)):
            texts.append(f"--- Page {i+1} ---\n{doc[i].get_text()[:8000]}")
        doc.close()
        out = "\n".join(texts)
        return out[:30000] + ("… (truncated)" if len(out) > 30000 else "")
    except ImportError:
        # Fallback to pypdf if PyMuPDF not installed
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            max_pages = max(1, min(int(max_pages), 50))
            texts = []
            for i in range(min(len(reader.pages), max_pages)):
                texts.append(f"--- Page {i+1} ---\n{reader.pages[i].extract_text() or ''[:8000]}")
            out = "\n".join(texts)
            return out[:30000] + ("… (truncated)" if len(out) > 30000 else "")
        except ImportError:
            return "Error: PDF support not installed. Install PyMuPDF (pip install pymupdf) or pypdf."
        except Exception as e:
            return f"Error reading PDF: {str(e)[:200]}"
    except Exception as e:
        return f"Error reading PDF: {str(e)[:200]}"

def docx_read(path: str) -> str:
    p, err = _guard_path(path)
    if err:
        return err
    if p.suffix.lower() != ".docx":
        return "Error: not a DOCX file."
    try:
        import docx
        doc = docx.Document(str(p))
        paras = [para.text for para in doc.paragraphs if para.text.strip()]
        tables = []
        for tbl in doc.tables:
            for row in tbl.rows:
                tables.append(" | ".join(cell.text.strip() for cell in row.cells))
        out = "\n".join(paras)
        if tables:
            out += "\n\n[Tables]\n" + "\n".join(tables)
        return out[:30000] + ("… (truncated)" if len(out) > 30000 else "")
    except ImportError:
        return "Error: python-docx not installed (pip install python-docx)."
    except Exception as e:
        return f"Error reading DOCX: {str(e)[:200]}"

def csv_analyze(path: str) -> str:
    p, err = _guard_path(path)
    if err:
        return err
    if p.suffix.lower() != ".csv":
        return "Error: not a CSV file."
    try:
        import csv
        with p.open(newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            if not rows:
                return "CSV is empty or has no data rows."
            cols = reader.fieldnames or []
            # Sample stats for numeric columns
            preview = rows[:5]
            lines = [f"CSV: {p.name} — {len(rows)} rows, {len(cols)} columns"]
            lines.append(f"Columns: {', '.join(cols)}")
            lines.append("\nPreview (first 5 rows):")
            for r in preview:
                lines.append(str({k: str(v)[:60] for k, v in r.items()}))
            # Numeric stats
            try:
                import statistics
                numeric_cols = []
                for c in cols:
                    vals = []
                    for r in rows:
                        try:
                            vals.append(float(r[c]))
                        except (ValueError, TypeError):
                            pass
                    if len(vals) > len(rows) * 0.5:
                        numeric_cols.append((c, vals))
                if numeric_cols:
                    lines.append("\nNumeric summary:")
                    for c, vals in numeric_cols:
                        lines.append(f"  {c}: min {min(vals):.2f} | max {max(vals):.2f} | mean {statistics.mean(vals):.2f}")
            except Exception:
                pass
            out = "\n".join(lines)
            return out[:8000]
    except Exception as e:
        return f"Error analyzing CSV: {str(e)[:200]}"

def register(registry: Any) -> None:
    registry.register(Tool(
        name="pdf_read",
        description="Extract text from a PDF file (local, secure, max 10MB, up to 50 pages).",
        parameters={
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Path to PDF file"},
                "max_pages": {"type": "integer", "description": "Max pages to extract (default 20)"},
            },
            "required": ["path"],
        },
        execute=pdf_read,
    ))
    registry.register(Tool(
        name="docx_read",
        description="Extract text and tables from a DOCX file (local, secure).",
        parameters={
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Path to DOCX file"},
            },
            "required": ["path"],
        },
        execute=docx_read,
    ))
    registry.register(Tool(
        name="csv_analyze",
        description="Analyze a CSV file: columns, preview, and numeric stats (local, secure).",
        parameters={
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Path to CSV file"},
            },
            "required": ["path"],
        },
        execute=csv_analyze,
    ))
